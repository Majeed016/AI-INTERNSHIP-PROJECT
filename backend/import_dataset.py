
import json
import re
from pathlib import Path
from collections import Counter

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "Summer Project_ Question Dataset.xlsx"
OUTPUT_FILE = BASE_DIR / "imported_questions.json"


def clean(value):
    return "" if value is None else str(value).strip()


def get_subject(sheet_name, row, headers):
    # Prefer the workbook's subject column when it contains a usable value.
    for i, header in enumerate(headers):
        if "subject" in clean(header).lower() and i < len(row):
            value = clean(row[i]).lower()
            if "math" in value:
                return "maths"
            if any(x in value for x in ("science", "physics", "chemistry", "biology")):
                return "science"

    name = sheet_name.lower()

    if "math" in name:
        return "maths"

    if any(x in name for x in ("science", "biology", "physics", "chemistry")):
        return "science"

    return None


def get_topic(sheet_name, subject):
    name = sheet_name.lower()

    for topic in ("algebra", "geometry", "trigonometry", "statistics"):
        if topic in name:
            return topic.title()

    for topic in ("physics", "chemistry", "biology"):
        if topic in name:
            return topic.title()

    # Preserve uncertainty instead of pretending every Maths question
    # belongs to Algebra or every Science question belongs to Physics.
    return "Unclassified"


def get_difficulty(cognitive_level):
    level = cognitive_level.lower()

    if "knowledge" in level or "remember" in level:
        return "easy"
    if "comprehension" in level or "understand" in level:
        return "medium"
    if any(word in level for word in ("application", "analysis", "evaluate", "create")):
        return "hard"

    return "unclassified"


def get_type(label):
    # Use only explicit question-type labels.
    # Do not mistake Bloom's cognitive levels for question types.
    label = label.strip().lower()

    if label == "mcq":
        return "mcq"
    if label == "short":
        return "short_answer"
    if label == "long":
        return "long_answer"

    return "unclassified"



def extract_single_mcq(text):
    # Accept only a simple question with exactly four clearly labelled options.
    pattern = re.compile(
        r"(?im)^\s*(?:A|B|C|D)[.)]\s*.+$"
    )
    options = [line.strip() for line in text.splitlines() if pattern.match(line)]

    if len(options) != 4:
        return None

    # Avoid treating a block containing multiple subquestions as one MCQ.
    if re.search(r"(?im)^\s*\(?[ivx]+\)\s+", text):
        return None

    cleaned_options = [
        re.sub(r"^\s*[A-D][.)]\s*", "", option)
        for option in options
    ]

    return cleaned_options


def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Workbook not found: {EXCEL_FILE}\n"
            "Copy the Excel workbook into the backend folder."
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True
    )

    imported = []
    seen = set()
    skipped = 0

    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            continue

        header_map = {
            clean(value).lower(): i
            for i, value in enumerate(headers)
            if value is not None
        }

        question_col = header_map.get("question")
        if question_col is None:
            continue

        level_col = header_map.get("cognitive level")
        type_col = header_map.get("short/long/mcq")

        for row_number, row in enumerate(rows, start=2):
            if question_col >= len(row):
                continue

            text = clean(row[question_col])
            if len(text) < 10:
                skipped += 1
                continue

            subject = get_subject(sheet.title, row, headers)
            if not subject:
                skipped += 1
                continue

            key = (subject, text.casefold())
            if key in seen:
                continue
            seen.add(key)

            level = (
                clean(row[level_col])
                if level_col is not None and level_col < len(row)
                else ""
            )
            label = (
                clean(row[type_col])
                if type_col is not None and type_col < len(row)
                else ""
            )

            question_type = get_type(label)
            options = None

            if question_type == "mcq":
                options = extract_single_mcq(text)

                # A grouped or malformed MCQ must be reviewed, not
                # passed to the generator as a valid single MCQ.
                if options is None:
                    question_type = "unclassified"

            marks = {
                "mcq": 2,
                "short_answer": 3,
                "long_answer": 6,
            }.get(question_type)

            question = {
                "id": f"IMP{len(imported) + 1:05d}",
                "subject": subject,
                "topic": get_topic(sheet.title, subject),
                "difficulty": get_difficulty(level),
                "type": question_type,
                "marks": marks,
                "text": text,
                "answer": "",
                "source": sheet.title,
                "source_row": row_number,
                "metadata_needs_review": (
                    question_type == "unclassified"
                    or question["topic"] == "Unclassified"
                    or question["difficulty"] == "unclassified"
                ),
                "marks_are_estimated": marks is not None,
            }

            if options is not None and question_type == "mcq":
                question["options"] = options

            imported.append(question)

    OUTPUT_FILE.write_text(
        json.dumps(imported, indent=2, ensure_ascii=False),
        encoding="utf-8"
    )

    print("Import complete!")
    print(f"Questions imported: {len(imported)}")
    print(f"Rows skipped: {skipped}")
    print("Question types:", dict(Counter(q["type"] for q in imported)))
    print("Subjects:", dict(Counter(q["subject"] for q in imported)))
    print(f"Output: {OUTPUT_FILE}")
    print("Existing question_bank.json was not changed.")


if __name__ == "__main__":
    main()
